import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_implementation_plan_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Keynote Dark Luxury Theme
    BG = RGBColor(11, 15, 25)           # #0B0F19 Deep Dark Navy
    CARD_BG = RGBColor(22, 30, 46)        # #161E2E Slate Card
    CARD_BORDER = RGBColor(45, 60, 88)    # #2D3C58 Card Border
    CYAN = RGBColor(6, 182, 212)          # #06B6D4 Electric Cyan
    RED = RGBColor(239, 68, 68)           # #EF4444 Crimson Alert
    GREEN = RGBColor(34, 197, 94)         # #22C55E Emerald Green
    WHITE = RGBColor(248, 250, 252)       # Crisp White
    MUTED = RGBColor(148, 163, 184)     # Slate Gray

    blank_layout = prs.slide_layouts[6]
    assets_dir = r"c:\cloud project\docs\assets"

    def set_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG
        bg.line.fill.background()

    def add_top_bar(slide, tag_text, title_text):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True

        p_tag = tf.paragraphs[0]
        p_tag.text = tag_text.upper()
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = CYAN
        p_tag.font.name = "Arial"

        p_ttl = tf.add_paragraph()
        p_ttl.text = title_text
        p_ttl.font.size = Pt(24)
        p_ttl.font.bold = True
        p_ttl.font.color.rgb = WHITE
        p_ttl.font.name = "Arial"
        p_ttl.space_before = Pt(2)

    def draw_card(slide, left, top, width, height, border_color=CARD_BORDER, bg_color=CARD_BG):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()
        return shape

    def add_fitted_image(slide, img_name, left, top, width, height, border_color=CYAN):
        draw_card(slide, left, top, width, height, border_color=border_color, bg_color=CARD_BG)
        img_path = os.path.join(assets_dir, img_name)
        if os.path.exists(img_path):
            pad = Inches(0.1)
            return slide.shapes.add_picture(img_path, left + pad, top + pad, width = width - (2*pad), height = height - (2*pad))
        return None

    # =========================================================================
    # SLIDE 1: Title Slide & Implementation Identity
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)

    tb1 = s1.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.733), Inches(1.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "CLOUD SECURITY ALERT LAMP 🚨💡"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = CYAN
    p.alignment = PP_ALIGN.CENTER

    p_sub = tf1.add_paragraph()
    p_sub.text = "S3 Security Scanner ➔ AWS IoT Core ➔ ESP32 Alert Lamp"
    p_sub.font.size = Pt(18)
    p_sub.font.bold = True
    p_sub.font.color.rgb = WHITE
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.space_before = Pt(8)

    authors = [
        ("D. Chander", "2503A52214", "Project Lead & Web Twin"),
        ("G. Rohan", "2503S52224", "System Architecture & Security"),
        ("K. Arjun", "2503A52235", "Backend & Database Developer"),
        ("P. Sachin", "2503A52240", "Hardware & Firmware Engineer")
    ]

    card_w = Inches(2.7)
    gap = Inches(0.25)
    left_start = Inches(0.8)

    for idx, (name, roll, role) in enumerate(authors):
        c_left = left_start + idx * (card_w + gap)
        draw_card(s1, c_left, Inches(2.7), card_w, Inches(3.8), border_color=CYAN)

        tb = s1.shapes.add_textbox(c_left + Inches(0.15), Inches(2.9), card_w - Inches(0.3), Inches(3.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p_name = tf.paragraphs[0]
        p_name.text = name
        p_name.font.size = Pt(17)
        p_name.font.bold = True
        p_name.font.color.rgb = WHITE

        p_roll = tf.add_paragraph()
        p_roll.text = f"Roll: {roll}"
        p_roll.font.size = Pt(12)
        p_roll.font.color.rgb = MUTED
        p_roll.space_before = Pt(4)

        p_role = tf.add_paragraph()
        p_role.text = role
        p_role.font.size = Pt(12)
        p_role.font.bold = True
        p_role.font.color.rgb = GREEN
        p_role.space_before = Pt(16)

    # =========================================================================
    # SLIDE 2: Project Overview & Compatibility Notes
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_top_bar(s2, "PROJECT OVERVIEW", "Cloud Security Alert Lamp & Hardware Pinouts")

    draw_card(s2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), border_color=CYAN)
    tb = s2.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "System Overview"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = CYAN

    points2 = [
        ("IoT-Enabled Security Auditor", "AWS Lambda security scanner audits S3 buckets and publishes alerts to AWS IoT Core over MQTT."),
        ("Sub-Second Desk Actuator", "ESP32 microcontroller receives MQTT alert over TLS and blinks Red LED rapidly to warn admins."),
        ("AWS Student / Learner Lab Ready", "Automated deployment supports standard AWS IAM roles and AWS Learner Lab pre-created LabRole."),
        ("Hardware Pinout Config", "Uses GPIO 4 for Red LED (with 220Ω resistor) and GPIO 2 for status LED (configurable in config.h).")
    ]
    for title, desc in points2:
        p = tf.add_paragraph()
        p.text = f"•  {title}"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.space_before = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = MUTED

    add_fitted_image(s2, "cloud_security_hero_lamp.png", Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), border_color=CYAN)

    # =========================================================================
    # SLIDE 3: Architecture Overview (Flowchart Diagram)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_top_bar(s3, "ARCHITECTURE OVERVIEW", "End-to-End Cyber-Physical Data Flow")

    add_fitted_image(s3, "cloud_pipeline_3d.png", Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), border_color=CYAN)

    draw_card(s3, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), border_color=GREEN)
    tb = s3.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Data Flow Execution"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = GREEN

    flow_steps = [
        ("S3 Event / Schedule", "EventBridge Rule or Cron triggers AWS Lambda Auditor Function."),
        ("3-Layer S3 Audit", "Lambda scans PublicAccessBlock, Bucket Policy & S3 ACL Grants."),
        ("IoT & SNS Dispatch", "Publishes payload to AWS IoT Core topic cloud/security/alerts & SNS."),
        ("Physical & Web Actuation", "ESP32 blinks Red LED via Secure TLS MQTT & Web Dashboard syncs.")
    ]
    for title, desc in flow_steps:
        p = tf.add_paragraph()
        p.text = f"•  {title}"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.space_before = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = MUTED

    # =========================================================================
    # SLIDE 4: Proposed Project Structure
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    add_top_bar(s4, "PROJECT STRUCTURE", "Repository Directories & Module Organization")

    modules = [
        ("cloud/", "AWS Backend & Auditor", CYAN, ["lambda/auditor.py (boto3 scanner)", "deploy_aws.py (1-click provisioner)", "test_vulnerable_bucket.py (CLI)", "aws_iot_policy.json (IAM schema)"]),
        ("firmware/", "ESP32 C++/Arduino Code", GREEN, ["src/main.cpp (non-blocking FSM)", "include/config.h (Wi-Fi & MQTT)", "include/secrets.h.template", "platformio.ini (build config)"]),
        ("web_dashboard/", "Web App & Simulator", CYAN, ["index.html (Dashboard & Simulator)", "styles.css (Glassmorphism UI)", "app.js (WebSocket visualizer)", "Interactive scenario runner"]),
        ("docs/", "Setup Guides & Wiring", RED, ["WIRING_GUIDE.md (Pinouts & Circuit)", "AWS_SETUP_GUIDE.md (Learner Lab)", "README.md (Master documentation)", "Step-by-step certificate guide"])
    ]

    for idx, (dir_name, role, color, files) in enumerate(modules):
        row = idx // 2
        col = idx % 2
        l = Inches(0.8) if col == 0 else Inches(6.8)
        t = Inches(1.8) if row == 0 else Inches(4.3)
        w = Inches(5.7)
        h = Inches(2.3)

        draw_card(s4, l, t, w, h, border_color=color)

        tb = s4.shapes.add_textbox(l + Inches(0.15), t + Inches(0.15), w - Inches(0.3), h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = dir_name
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = color

        p_sub = tf.add_paragraph()
        p_sub.text = role.upper()
        p_sub.font.size = Pt(9)
        p_sub.font.bold = True
        p_sub.font.color.rgb = WHITE
        p_sub.space_before = Pt(1)

        for f in files:
            p_f = tf.add_paragraph()
            p_f.text = f"•  {f}"
            p_f.font.size = Pt(9.5)
            p_f.font.color.rgb = MUTED
            p_f.space_before = Pt(1)

    # =========================================================================
    # SLIDE 5: AWS Cloud Security Backend (cloud/)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_top_bar(s5, "AWS CLOUD BACKEND", "Lambda Auditor, Deployer & Test CLI (cloud/)")

    add_fitted_image(s5, "security_layers_diag.png", Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), border_color=CYAN)

    draw_card(s5, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), border_color=RED)
    tb = s5.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Backend Component Details"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = RED

    be_details = [
        ("Lambda Auditor (auditor.py)", "Scans S3 GetPublicAccessBlock, GetBucketPolicyStatus & GetBucketAcl. Publishes JSON ALERT/OK payload."),
        ("1-Click Deployer (deploy_aws.py)", "Provisions IoT Thing, Policy, Certs, Lambda zip, EventBridge cron (5 mins), and SNS Topic in < 60s."),
        ("Test Utility (test_vulnerable_bucket.py)", "CLI tool: python test_vulnerable_bucket.py create (triggers lamp) and delete (resolves alert).")
    ]
    for title, desc in be_details:
        p = tf.add_paragraph()
        p.text = f"•  {title}"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.space_before = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = MUTED

    # =========================================================================
    # SLIDE 6: ESP32 Firmware Architecture (firmware/)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)
    add_top_bar(s6, "ESP32 FIRMWARE ENGINE", "TLS MQTT Client & FreeRTOS Non-Blocking State Machine")

    add_fitted_image(s6, "crypto_shield_vault.png", Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), border_color=CYAN)

    draw_card(s6, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), border_color=GREEN)
    tb = s6.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Firmware Architecture (firmware/)"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = GREEN

    fw_details = [
        ("Framework & Libraries", "Built on C++/Arduino using WiFiClientSecure, PubSubClient, and ArduinoJson."),
        ("TLS X.509 Authentication", "Handles mutual TLS authentication with AWS IoT Core over Port 8883."),
        ("Non-Blocking Millis FSM", "Blink state machine uses millis() timers without freezing delay() calls."),
        ("Auto-Reconnect Engine", "Monitors Wi-Fi and MQTT connection loss with automatic exponential reconnects.")
    ]
    for title, desc in fw_details:
        p = tf.add_paragraph()
        p.text = f"•  {title}"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.space_before = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = MUTED

    # =========================================================================
    # SLIDE 7: Web Dashboard & Simulator (web_dashboard/)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7)
    add_top_bar(s7, "WEB DASHBOARD & SIMULATOR", "Real-Time S3 Monitor & Hardware Lamp Visualizer")

    draw_card(s7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), border_color=CYAN)
    tb = s7.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Web Dashboard Capabilities"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = CYAN

    web_details = [
        ("Real-Time S3 Monitor", "Renders live security status, exposed bucket names, and scan timestamps."),
        ("Hardware Lamp Simulator", "Simulated glowing 2D/3D Red Alert Lamp with realistic blinking animation."),
        ("Test Scenario Runner", "Interactive browser buttons to simulate bucket vulnerability creation and fix actions."),
        ("Learner Lab Wizard", "Step-by-step interactive manual for AWS Student Account certificate configuration.")
    ]
    for title, desc in web_details:
        p = tf.add_paragraph()
        p.text = f"•  {title}"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.space_before = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = MUTED

    add_fitted_image(s7, "neon_dashboard_screen.png", Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), border_color=CYAN)

    # =========================================================================
    # SLIDE 8: Setup, Wiring & Verification Plan
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8)
    add_top_bar(s8, "VERIFICATION & WIRING PLAN", "ESP32 Circuit Schematics & End-to-End Verification")

    draw_card(s8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), border_color=CYAN)
    tb = s8.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Hardware Wiring (docs/)"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = CYAN

    hw_plan = [
        ("Circuit Wiring Schematics", "ESP32 DevKit v1 + High-Brightness Red LED + 220Ω limiting resistor on GPIO 4."),
        ("AWS Student Guide", "Tailored step-by-step setup guide for $50 AWS Learner Lab accounts.")
    ]
    for title, desc in hw_plan:
        p = tf.add_paragraph()
        p.text = f"•  {title}"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.space_before = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = MUTED

    draw_card(s8, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), border_color=GREEN)
    tb = s8.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Verification Plan"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = GREEN

    ver_plan = [
        ("Automated CLI Test", "Run test_vulnerable_bucket.py create/delete to verify Lambda & IoT Core payload delivery."),
        ("Hardware Verification", "Verify TLS payload parsing and state transition (OK ➔ ALERT 200ms blink ➔ OK)."),
        ("Web Simulator Test", "Verify live blinking state animations and dashboard controls in browser.")
    ]
    for title, desc in ver_plan:
        p = tf.add_paragraph()
        p.text = f"•  {title}"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.space_before = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = MUTED

    # Save presentation
    out_paths = [
        "docs/Cloud_Security_Alert_Lamp_Presentation.pptx",
        "docs/Cloud_Security_Alert_Lamp_Minimal.pptx",
        "docs/Cloud_Security_Alert_Lamp_Executive.pptx"
    ]
    for pth in out_paths:
        try:
            prs.save(pth)
            print(f"[+] Successfully saved Implementation Plan Deck to: {pth}")
        except PermissionError:
            alt = pth.replace(".pptx", "_plan.pptx")
            prs.save(alt)
            print(f"[+] Saved to plan path: {alt}")

if __name__ == "__main__":
    create_implementation_plan_deck()
